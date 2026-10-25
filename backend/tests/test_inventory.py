"""Inventory, resource detail, orphans, export, changes timeline and compare (demo tenant)."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import psycopg
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from kloudrecon.services.azure.demo.world import world_at
from kloudrecon.services.scan.pipeline import run_scan
from tests.conftest import auth_headers
from tests.test_scans import _demo_org, new_scan

pytestmark = pytest.mark.usefixtures("seeded")

VM = "microsoft.compute/virtualmachines"


@pytest.fixture
def demo(admin_conn: psycopg.Connection[tuple[object, ...]], client: TestClient) -> dict[str, Any]:
    org_id = _demo_org()
    conn_id = str(
        admin_conn.execute("SELECT id FROM connections WHERE org_id = %s", (org_id,)).fetchone()[0]  # type: ignore[index]
    )
    now = datetime.now(UTC)
    first = new_scan(conn_id, "seed")
    assert run_scan(first, as_of=now - timedelta(days=8)) == "succeeded"
    second = new_scan(conn_id, "seed")
    assert run_scan(second, as_of=now - timedelta(hours=2)) == "succeeded"
    admin_conn.execute(
        "INSERT INTO users (org_id, external_oid, email, name, org_role, status) "
        "VALUES (%s, 'dev:demo@contoso.test', 'demo@contoso.test', 'Demo', 'admin', 'active')",
        (org_id,),
    )
    return {
        "headers": auth_headers(client, "demo@contoso.test"),
        "first": str(first),
        "second": str(second),
        "conn": conn_id,
    }


def get(client: TestClient, demo: dict[str, Any], path: str, **params: Any) -> Any:
    resp = client.get(path, params=params, headers=demo["headers"])
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_resources_latest_paging_and_time_machine(client: TestClient, demo: dict[str, Any]) -> None:
    page = get(client, demo, "/api/v1/resources", limit=250)
    total = page["total"]
    assert total == len(world_at(1).resources)
    assert page["scan_ids"] == [demo["second"]]
    seen = {r["id"] for r in page["items"]}
    cursor = page["next_cursor"]
    while cursor:
        page = get(client, demo, "/api/v1/resources", limit=250, cursor=cursor)
        seen |= {r["id"] for r in page["items"]}
        cursor = page["next_cursor"]
    assert len(seen) == total
    older = get(client, demo, "/api/v1/resources", scan_id=demo["first"], limit=1)
    assert older["total"] == len(world_at(0).resources)


def test_resource_filters_and_sort(client: TestClient, demo: dict[str, Any]) -> None:
    vms = get(client, demo, "/api/v1/resources", type=VM, limit=500)
    expected = sum(1 for r in world_at(1).resources.values() if r["type"] == VM)
    assert vms["total"] == expected
    assert {r["type"] for r in vms["items"]} == {VM}
    jump = get(client, demo, "/api/v1/resources", q="jumpbox", type=VM)
    assert [r["name"] for r in jump["items"]] == ["vm-prod-jumpbox01"]
    assert jump["items"][0]["sku"] == "Standard_B2ms"
    assert jump["items"][0]["subscription_name"] == "Contoso-Prod"
    dev = get(client, demo, "/api/v1/resources", tag="env=dev", limit=500)
    assert dev["total"] > 50
    assert all(r["tags"].get("env") == "dev" for r in dev["items"])
    by_type = get(client, demo, "/api/v1/resources", sort="type", limit=50)
    types = [r["type"] for r in by_type["items"]]
    assert types == sorted(types)


def test_facets(client: TestClient, demo: dict[str, Any]) -> None:
    facets = get(client, demo, "/api/v1/resources/facets")
    types = {f["value"]: f["count"] for f in facets["types"]}
    assert types[VM] > 60
    assert {f["label"] for f in facets["subscriptions"]} >= {"Contoso-Prod", "Contoso-NonProd"}
    assert "env" in {f["value"] for f in facets["tag_keys"]}


def test_resource_detail(client: TestClient, demo: dict[str, Any]) -> None:
    jump = get(client, demo, "/api/v1/resources", q="vm-prod-jumpbox01", type=VM)["items"][0]
    detail = get(client, demo, "/api/v1/resources/detail", arm_id=jump["arm_id"])
    assert detail["properties"]["hardwareProfile"]["vmSize"] == "Standard_B2ms"
    related = {(r["edge_type"], r["name"]) for r in detail["related"]}
    assert ("nic", "nic-vm-prod-jumpbox01") in related
    assert ("disk", "osdisk-vm-prod-jumpbox01") in related
    assert detail["portal_url"].startswith("https://portal.azure.com/#@")
    assert detail["portal_url"].endswith(jump["arm_id"])
    resp = client.get(
        "/api/v1/resources/detail",
        params={"arm_id": "/subscriptions/x/nope"},
        headers=demo["headers"],
    )
    assert resp.status_code == 404


def test_orphans(client: TestClient, demo: dict[str, Any]) -> None:
    result = get(client, demo, "/api/v1/resources/orphans")
    cats = result["by_category"]
    for category in (
        "unattached_disk",
        "unused_public_ip",
        "unattached_nic",
        "empty_app_service_plan",
        "old_snapshot",
        "empty_resource_group",
        "deallocated_vm",
        "unused_nsg",
        "empty_load_balancer",
        "unused_route_table",
    ):
        assert cats.get(category, 0) >= 1, category
    names = {o["name"] for o in result["items"]}
    assert "asp-dev-unused-p1v3" in names
    assert "rg-empty-temp" in names
    assert "kubernetes-internal" not in names  # AKS-managed LB is not an orphan


def test_export_csv_and_xlsx(
    client: TestClient, demo: dict[str, Any], admin_conn: psycopg.Connection[tuple[object, ...]]
) -> None:
    total = get(client, demo, "/api/v1/resources", type=VM, limit=1)["total"]
    csv_resp = client.get(
        "/api/v1/resources/export", params={"format": "csv", "type": VM}, headers=demo["headers"]
    )
    assert csv_resp.status_code == 200
    assert "attachment" in csv_resp.headers["content-disposition"]
    lines = csv_resp.content.decode("utf-8-sig").strip().splitlines()
    assert lines[0].startswith("name,type,resource_group")
    assert len(lines) == total + 1
    xlsx = client.get(
        "/api/v1/resources/export", params={"format": "xlsx"}, headers=demo["headers"]
    )
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content), read_only=True)["Inventory"]
    assert sum(1 for _ in ws.iter_rows()) == len(world_at(1).resources) + 1
    actions = [r[0] for r in admin_conn.execute("SELECT action FROM audit_logs").fetchall()]
    assert actions.count("inventory.exported") == 2


def test_export_neutralizes_formulas() -> None:
    from kloudrecon.services.inventory import to_csv

    out = to_csv(iter([["name"], ["=HYPERLINK(1)"], ["+cmd"], ["safe"]])).decode("utf-8-sig")
    assert "'=HYPERLINK(1)" in out
    assert "'+cmd" in out


def test_changes_timeline_filters_and_paging(client: TestClient, demo: dict[str, Any]) -> None:
    created = get(client, demo, "/api/v1/changes", change_type="created", days=30)
    assert created["items"]
    assert {c["change_type"] for c in created["items"]} == {"created"}
    assert any(c["caller"] for c in created["items"])
    risky = get(client, demo, "/api/v1/changes", risky=True, days=90)
    assert risky["items"]
    assert all(c["risky"] and c["risk_reason"] for c in risky["items"])
    page = get(client, demo, "/api/v1/changes", days=90, limit=40)
    ids = [c["id"] for c in page["items"]]
    while page["next_cursor"]:
        page = get(client, demo, "/api/v1/changes", days=90, limit=40, cursor=page["next_cursor"])
        ids += [c["id"] for c in page["items"]]
    assert len(ids) == len(set(ids)) > 100
    times = [c["occurred_at"] for c in get(client, demo, "/api/v1/changes", days=90)["items"]]
    assert times == sorted(times, reverse=True)
    summary = get(client, demo, "/api/v1/changes/summary", days=30)
    assert len(summary["days"]) == 30
    assert summary["totals"].get("created", 0) >= 3


def test_compare_two_scans(client: TestClient, demo: dict[str, Any]) -> None:
    result = get(client, demo, "/api/v1/compare", **{"from": demo["first"], "to": demo["second"]})
    counts = result["counts"]
    assert counts["from_total"] == len(world_at(0).resources)
    assert counts["to_total"] == len(world_at(1).resources)
    assert counts["added"] >= 3
    assert {i["name"] for i in result["added"]} >= {"vm-dev-feature01", "nic-vm-dev-feature01"}
    resized = [
        i
        for i in result["changed"]
        if any(f["path"] == "properties.hardwareProfile.vmSize" for f in i["fields"])
    ]
    assert resized
    assert any(t["type"] == VM for t in result["by_type"])
    swapped = get(client, demo, "/api/v1/compare", **{"from": demo["second"], "to": demo["first"]})
    assert swapped["counts"] == counts
    missing = client.get(
        "/api/v1/compare",
        params={"from": str(uuid.uuid4()), "to": demo["second"]},
        headers=demo["headers"],
    )
    assert missing.status_code == 404


def test_inventory_is_tenant_isolated(client: TestClient, demo: dict[str, Any]) -> None:
    acme = auth_headers(client, "admin@acme.test")
    assert client.get("/api/v1/resources", headers=acme).json()["total"] == 0
    resp = client.get("/api/v1/resources", params={"scan_id": demo["second"]}, headers=acme)
    assert resp.status_code == 404
    resp = client.get(
        "/api/v1/compare", params={"from": demo["first"], "to": demo["second"]}, headers=acme
    )
    assert resp.status_code == 404
    assert client.get("/api/v1/changes", headers=acme).json()["items"] == []
    assert client.get("/api/v1/resources/orphans", headers=acme).json()["items"] == []


def test_overview(client: TestClient, demo: dict[str, Any]) -> None:
    data = get(client, demo, "/api/v1/overview")
    assert data["counts"]["resources"] == len(world_at(1).resources)
    assert data["counts"]["subscriptions"] == 4
    assert data["counts"]["orphans"] > 10
    assert data["latest_scan"]["id"] == demo["second"]
    assert data["active_scan"] is None
    assert data["changes_since_last_scan"]["created"] >= 3
    assert len(data["history"]) == 2
    assert data["by_type"][0]["count"] >= data["by_type"][-1]["count"]
    links = {a["link"] for a in data["attention"]}
    assert "/app/inventory?tab=orphans" in links
    assert data["recent_changes"]
    older = get(client, demo, "/api/v1/overview", scan_id=demo["first"])
    assert older["counts"]["resources"] == len(world_at(0).resources)


def test_overview_without_connections(client: TestClient) -> None:
    data = client.get("/api/v1/overview", headers=auth_headers(client, "admin@acme.test")).json()
    assert data["connections"] == []
    assert data["counts"]["resources"] == 0
    assert data["latest_scan"] is None
