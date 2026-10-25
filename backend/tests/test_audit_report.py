"""The Word writer of the audit report (decision 0053): cover, approval page, contents,
front matter and blocks, without a browser. The demo-tenant run is in test_reports.py."""

from __future__ import annotations

import io
from typing import Any

from docx import Document

from kloudrecon.reports.document import Branding, ReportDoc
from kloudrecon.reports.word import Control, render_docx


def _doc(data: bytes) -> Any:
    return Document(io.BytesIO(data))


def _text(d: Any) -> str:
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_word_writer_front_matter_and_blocks() -> None:
    src = ReportDoc(
        title="Azure Audit Report",
        subtitle="Security and cost",
        branding=Branding(company="Contoso Ltd", watermark="TRIAL"),
        meta=[
            ("Prepared for", "Contoso Ltd"),
            ("Date", "04 October 2026"),
            ("Scope", "2 subscriptions"),
        ],
    )
    src.section("Executive summary", "Lead paragraph.")
    src.kpis([("Resources", "57", "2 subscriptions"), ("Critical risks", "3", None)])
    src.sub("Highest risks")
    src.table(
        ["Severity", "Finding", "Risk"],
        [("critical", "RDP open", 92), ("low", "Tag missing", 10)],
        numeric=[2],
        badges={0: {"critical": "critical"}},
    )
    src.figure("<svg></svg>", "Monthly cost")
    src.callout("Mapped, not certified", "Alignment only.")
    src.points([("Close RDP", "Use Bastion")], tone="critical")
    src.bullets(["one", "two"])
    src.section("Financial summary")
    src.p("Spend after savings.")
    data = render_docx(src, Control(approved_by="Dana Reyes", scan="#4"))
    assert data[:2] == b"PK"
    d = _doc(data)
    text = _text(d)
    # page 1: text cover (no browser), page 2: document control and approval
    assert "Azure Audit Report" in text
    assert "Document control" in text
    for label in ("Prepared by", "Reviewed by", "Report approved by", "Signature", "Distribution"):
        assert label in text
    assert "Dana Reyes" in text
    assert "Confidentiality." in text
    # contents: a TOC field whose cached result lists the sections
    xml = d.element.xml
    assert r'TOC \o "1-2"' in xml
    assert "updateFields" in d.settings.element.xml
    headings = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["1  Executive summary", "2  Financial summary"]
    assert [p.text for p in d.paragraphs if p.style.name == "Heading 2"] == ["1.1  Highest risks"]
    # blocks
    assert "RDP open" in text
    assert "Figure (not rendered): Monthly cost" in text
    assert "Mapped, not certified" in text
    assert "Close RDP" in text
    # header and footer on every page but the cover
    sec = d.sections[0]
    assert sec.different_first_page_header_footer is True
    assert "Contoso Ltd" in sec.header.paragraphs[0].text
    assert "TRIAL" in sec.header.paragraphs[0].text
    assert "Confidential" in sec.footer.paragraphs[0].text
    assert "NUMPAGES" in sec.footer._element.xml
    assert d.core_properties.title == "Azure Audit Report"


def test_cover_picture_is_a_full_page_picture_behind_the_text() -> None:
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    src = ReportDoc(title="T", subtitle="S", branding=Branding(company="C"), meta=[])
    src.section("One")
    src.figure("<svg></svg>", "A chart")
    d = _doc(render_docx(src, Control(), cover=png, figures={1: png}))
    xml = d.element.xml
    assert "wp:anchor" in xml
    assert 'behindDoc="1"' in xml
    assert "Figure: A chart" in _text(d)
    assert len(d.inline_shapes) == 1  # the figure; the cover is anchored, not inline
