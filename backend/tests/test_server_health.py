"""Server health checks (decision 0054): the analysis of script output (pure) and the API
(servers, script, pasted output, tenant isolation, the Word report)."""

# ruff: noqa: E501, S104 -- sample script outputs (long lines, 0.0.0.0 listeners)

from __future__ import annotations

import io
import json
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from kloudrecon.health import analyze as engine
from kloudrecon.health import catalog
from tests.conftest import auth_headers, register, run_jobs

NOW = datetime(2026, 10, 4, 12, 0, tzinfo=UTC)
SCRIPTS = Path(engine.__file__).parent / "scripts"


def windows(**over: Any) -> dict[str, Any]:
    data: dict[str, Any] = {
        "schema": "kloudrecon.windows.v1",
        "server": "srv-app01",
        "collected_at": "2026-10-04T10:00:00Z",
        "elevated": True,
        "os": {
            "caption": "Microsoft Windows Server 2016 Datacenter",
            "version": "10.0.14393",
            "build": "14393",
            "architecture": "64-bit",
            "hostname": "SRV-APP01",
            "domain": "corp.example",
            "last_boot": (NOW - timedelta(days=400)).isoformat(),
            "manufacturer": "Microsoft Corporation",
            "model": "Virtual Machine",
        },
        "cpu": {"name": "Xeon", "cores": 4, "logical": 8, "load_percent": 35},
        "memory": {
            "total_bytes": 16 * 2**30,
            "free_bytes": 1 * 2**29,
            "pagefile_mb": 4096,
            "pagefile_used_mb": 500,
        },
        "disks": [
            {
                "name": "C:",
                "filesystem": "NTFS",
                "size_bytes": 100 * 2**30,
                "free_bytes": 3 * 2**30,
            },
            {
                "name": "D:",
                "filesystem": "NTFS",
                "size_bytes": 500 * 2**30,
                "free_bytes": 300 * 2**30,
            },
        ],
        "updates_checked": True,
        "pending_updates": [
            {
                "title": "2026-09 Cumulative Update (KB5043080)",
                "kb": "5043080",
                "severity": "Critical",
                "categories": "Security Updates",
                "cves": ["CVE-2026-1111", "CVE-2026-2222"],
                "published": "2026-09-10",
            },
            {
                "title": "Security Update for .NET (KB5042352)",
                "kb": "5042352",
                "severity": "Important",
                "categories": "Security Updates",
                "cves": [],
                "published": "2026-08-12",
            },
            {
                "title": "Security Intelligence Update for Defender",
                "kb": "2267602",
                "severity": "",
                "categories": "Definition Updates",
                "cves": [],
            },
        ],
        "hotfixes": [
            {"id": "KB5031111", "description": "Security Update", "installed_on": "2026-02-10"}
        ],
        "reboot_pending": True,
        "services_stopped": [
            {
                "name": "MSSQLSERVER",
                "display_name": "SQL Server (MSSQLSERVER)",
                "state": "Stopped",
                "delayed": False,
            },
            {
                "name": "edgeupdate",
                "display_name": "Microsoft Edge Update",
                "state": "Stopped",
                "delayed": True,
            },
        ],
        "event_counts": {
            "System": {"critical": 2, "error": 40, "warning": 10},
            "Application": {"critical": 0, "error": 5, "warning": 3},
        },
        "events": [
            {
                "log": "System",
                "source": "disk",
                "id": 7,
                "level": "error",
                "count": 12,
                "last": "2026-10-03T08:00:00Z",
                "message": "The device has a bad block.",
            }
        ],
        "firewall": [
            {"profile": "Domain", "enabled": True},
            {"profile": "Public", "enabled": False},
        ],
        "defender": {
            "available": True,
            "service_enabled": True,
            "realtime": False,
            "signature_age_days": 12,
        },
        "smb1_enabled": True,
        "rdp": {"enabled": True, "nla": False},
        "tls": {"SSL 3.0": None, "TLS 1.0": 1, "TLS 1.1": None, "TLS 1.2": 1},
        "local_admins": [{"name": f"CORP\\admin{i}", "type": "User"} for i in range(7)],
        "listening": [
            {"port": 3389, "address": "0.0.0.0", "process": "svchost"},
            {"port": 23, "address": "0.0.0.0", "process": "tlntsvr"},
        ],
        "certificates_expiring": [
            {"subject": "CN=app.corp.example", "not_after": "2026-10-20T00:00:00Z"}
        ],
    }
    data.update(over)
    return data


def linux(**sections: str) -> dict[str, Any]:
    base = {
        "os_release": 'PRETTY_NAME="Ubuntu 20.04.6 LTS"\nNAME="Ubuntu"\nVERSION_ID="20.04"\nID=ubuntu',
        "hostname": "web01.corp.example",
        "uname": "Linux 5.4.0-190-generic x86_64",
        "uptime": "864000.10 100.0",
        "cpu": "cores 4\nmodel name\t: Intel Xeon\nloadavg 9.50 3.00 2.00 1/200 999",
        "meminfo": "MemTotal:        8000000 kB\nMemFree:          200000 kB\nMemAvailable:     300000 kB\nSwapTotal:       2000000 kB\nSwapFree:         200000 kB",
        "df": "Filesystem Type 1024-blocks Used Available Capacity Mounted on\n/dev/sda1 ext4 100000000 97000000 3000000 97% /\n/dev/sdb1 xfs 500000000 100000000 400000000 20% /data",
        "inodes": "Filesystem Inodes IUsed IFree IUse% Mounted on\n/dev/sda1 1000 950 50 95% /\n/dev/sdb1 1000 10 990 1% /data",
        "virtualization": "microsoft",
        "failed_units": "nginx.service failed\nbackup.service failed",
        "journal_errors": "2026-10-03T10:00:00+0000 web01 nginx[1234]: upstream timed out\n2026-10-03T11:00:00+0000 web01 nginx[1300]: upstream timed out\n2026-10-03T12:00:00+0000 web01 kernel: EXT4-fs error (device sda1)",
        "journal_counts": "2 1\n3 620\n4 80",
        "kernel_log": "",
        "oom_kills": "2",
        "reboot_required": "yes\nlinux-image-5.4.0-200-generic",
        "kernels": "running 5.4.0-190-generic\ninstalled 5.4.0-200-generic",
        "sshd": "port 22\npermitrootlogin yes\npasswordauthentication yes\npermitemptypasswords no",
        "listening": "tcp 0.0.0.0:22\ntcp 0.0.0.0:21\ntcp 127.0.0.1:5432",
        "firewall": "ufw Status: inactive\niptables_rules 0",
        "mac": "apparmor enabled",
        "time": "NTPSynchronized=no\nTimezone=UTC",
        "accounts": "uid0 root\nuid0 toor\nadmins alice,bob\nempty_passwords 1\nlogin_users 4",
        "patch_state": "manager apt\nlists_updated 1790000000\nlast_install 2026-03-01\nauto_updates 0\npackages 640",
        "patches": "openssl/focal-updates,focal-security 1.1.1f-1ubuntu2.24 amd64 [upgradable from: 1.1.1f-1ubuntu2.20]\nvim/focal-updates 2:8.1 amd64 [upgradable from: 2:8.0]",
        "updates_count": "2",
        "cves": "",
    }
    base.update(sections)
    return {
        "schema": "kloudrecon.linux.v1",
        "server": "web01",
        "collected_at": "2026-10-04T10:00:00Z",
        "root": True,
        "sections": base,
    }


def ids(summary: dict[str, Any]) -> dict[str, str]:
    return {f["id"]: f["severity"] for f in summary["findings"]}


# ---- parsing -------------------------------------------------------------------------------------
def test_parse_accepts_console_text_around_the_json_and_rejects_everything_else() -> None:
    text = (
        "PS C:\\> .\\check.ps1\n"
        + json.dumps(windows())
        + "\nDone. The result is on the clipboard."
    )
    platform, data = engine.parse_output(text)
    assert (platform, data["server"]) == ("windows", "srv-app01")
    for bad in ("hello", "{not json}", json.dumps({"schema": "other"}), json.dumps([1, 2])):
        with pytest.raises(engine.HealthInputError):
            engine.parse_output(bad)
    with pytest.raises(engine.HealthInputError, match="too large"):
        engine.parse_output("{" + " " * (engine.MAX_BYTES + 10) + "}")


# ---- Windows -------------------------------------------------------------------------------------
def test_windows_report_finds_capacity_patching_and_security_problems() -> None:
    _, _, s = engine.analyze(json.dumps(windows()), "windows-server", "2016", NOW)
    found = ids(s)
    assert found["disk-C:"] == "critical"  # 97% full
    assert "disk-D:" not in found
    assert found["memory"] == "high"  # 97% used
    assert found["uptime"] == "high"  # 400 days
    assert found["os-eol"] == "medium"  # Windows Server 2016 ends 2027-01-12
    assert found["patch-critical"] == "high"
    assert found["patch-age"] == "high"
    assert found["reboot"] == "medium"
    assert found["events-critical"] == "medium"
    assert found["events-disk"] == "high"
    assert found["services"] == "medium"
    assert found["firewall"] == "high"
    assert found["antivirus"] == "high"
    assert found["signatures"] == "medium"
    assert found["smb1"] == "high"
    assert found["rdp-nla"] == "high"
    assert found["tls"] == "low"
    assert found["admins"] == "low"
    assert found["certificates"] == "medium"
    assert found["legacy-ports"] == "medium"
    assert s["status"] == "critical"
    assert s["score"] == 0
    assert s["findings"][0]["severity"] == "critical"  # worst first
    p = s["patches"]
    assert (p["pending"], p["security"], p["critical"]) == (2, 2, 1)  # definitions are not counted
    assert p["items"][0]["id"] == "KB5043080"
    assert {v["id"] for v in s["vulnerabilities"]} == {"CVE-2026-1111", "CVE-2026-2222"}
    assert [x["name"] for x in s["services"]["failed"]] == ["MSSQLSERVER"]  # updaters are ignored
    assert s["os"]["lifecycle"]["end_of_support"] == "2027-01-12"
    assert all(f["recommendation"] for f in s["findings"])
    checks = {c["key"]: c["status"] for c in s["security"]}
    assert checks["firewall"] == "fail"
    assert checks["smb1"] == "fail"


def test_a_healthy_windows_server_scores_100() -> None:
    good = windows(
        os={
            **windows()["os"],
            "caption": "Microsoft Windows Server 2022 Datacenter",
            "last_boot": (NOW - timedelta(days=20)).isoformat(),
        },
        memory={
            "total_bytes": 16 * 2**30,
            "free_bytes": 8 * 2**30,
            "pagefile_mb": 4096,
            "pagefile_used_mb": 100,
        },
        disks=[
            {
                "name": "C:",
                "filesystem": "NTFS",
                "size_bytes": 100 * 2**30,
                "free_bytes": 60 * 2**30,
            }
        ],
        pending_updates=[],
        hotfixes=[{"id": "KB1", "description": "Security Update", "installed_on": "2026-09-20"}],
        reboot_pending=False,
        services_stopped=[],
        event_counts={"System": {"critical": 0, "error": 3, "warning": 1}},
        events=[],
        firewall=[{"profile": "Domain", "enabled": True}],
        defender={"available": True, "realtime": True, "signature_age_days": 0},
        smb1_enabled=False,
        rdp={"enabled": True, "nla": True},
        tls={},
        local_admins=[{"name": "CORP\\admins"}],
        listening=[{"port": 443, "address": "0.0.0.0", "process": "System"}],
        certificates_expiring=[],
    )
    _, _, s = engine.analyze(json.dumps(good), "windows-server", "2022", NOW)
    assert s["findings"] == []
    assert (s["score"], s["grade"], s["status"]) == (100, "A", "healthy")
    # without administrator rights the report says it is partial
    _, _, partial = engine.analyze(
        json.dumps({**good, "elevated": False, "updates_checked": False}),
        "windows-server",
        "2022",
        NOW,
    )
    assert ids(partial) == {"not-elevated": "low", "patch-unknown": "low"}
    assert partial["complete"] is False


# ---- Linux ---------------------------------------------------------------------------------------
def test_linux_report_from_apt_output() -> None:
    _, _, s = engine.analyze(json.dumps(linux()), "ubuntu", "20.04", NOW)
    found = ids(s)
    assert found["disk-/"] == "critical"
    assert found["inodes-/"] == "high"
    assert found["memory"] == "high"
    assert found["swap"] == "medium"
    assert found["cpu"] == "high"  # load 9.5 on 4 cores
    assert found["os-eol"] == "high"  # Ubuntu 20.04 ended 2025-05-31
    assert found["patch-security"] == "medium"
    assert found["patch-age"] == "high"
    assert found["reboot"] == "medium"
    assert found["auto-updates"] == "low"
    assert found["oom"] == "high"
    assert found["disk-errors"] == "high"
    assert found["events-errors"] == "medium"
    assert found["services"] == "medium"
    assert found["ssh-root"] == "high"
    assert found["ssh-password"] == "medium"
    assert found["firewall"] == "medium"
    assert found["uid0"] == "critical"
    assert found["empty-passwords"] == "critical"
    assert found["time"] == "low"
    assert found["legacy-ports"] == "medium"
    assert s["status"] == "critical"
    assert s["os"]["name"] == "Ubuntu 20.04.6 LTS"
    assert s["os"]["kernel"] == "5.4.0-190-generic"
    assert s["os"]["uptime_days"] == 10.0
    assert [d["mount"] for d in s["disks"]] == ["/", "/data"]
    p = s["patches"]
    assert (p["checked"], p["pending"], p["security"], p["last_patched"]) == (
        True,
        2,
        1,
        "2026-03-01",
    )
    assert p["items"][0] == {
        "name": "openssl",
        "id": "1.1.1f-1ubuntu2.24",
        "severity": "Security",
        "security": True,
        "cves": [],
        "published": "",
    }
    assert s["logs"]["top"][0]["source"] == "nginx"
    assert s["logs"]["top"][0]["count"] == 2
    assert {(x["port"], x["address"]) for x in s["network"]["listening"]} == {
        (22, "0.0.0.0"),
        (21, "0.0.0.0"),
        (5432, "127.0.0.1"),
    }


def test_linux_rpm_advisories_and_cves() -> None:
    data = linux(
        os_release='PRETTY_NAME="Rocky Linux 9.4"\nVERSION_ID="9.4"',
        patch_state="manager rpm\nlast_install Wed 19 Nov 2025 10:25:22 PM UTC\nauto_updates 1\npackages 400",
        patches="RLSA-2025:23343 Critical/Sec.  openssl-3.0.7-28.el9.x86_64\nRLSA-2025:1111 Moderate/Sec.  binutils-2.35.2-67.el9.x86_64\nRLBA-2024:9438  bugfix         bash-5.1.8-9.el9.x86_64",
        updates_count="3",
        cves="CVE-2025-11083  Moderate/Sec.  binutils-2.35.2-67.el9.x86_64\nCVE-2025-9999   Critical/Sec.  openssl-3.0.7-28.el9.x86_64",
        reboot_required="no",
    )
    _, _, s = engine.analyze(json.dumps(data), "rocky", "9", NOW)
    p = s["patches"]
    assert (p["pending"], p["security"], p["critical"]) == (3, 2, 1)
    assert p["last_patched"] == "2025-11-19"
    assert p["auto_updates"] is True
    assert ids(s)["patch-critical"] == "high"
    assert "auto-updates" not in ids(s)
    assert {v["id"] for v in s["vulnerabilities"]} == {"CVE-2025-11083", "CVE-2025-9999"}
    assert (
        "CVE-2025-9999" in next(f for f in s["findings"] if f["id"] == "patch-critical")["detail"]
    )


def test_linux_zypper_table_and_unreachable_repositories() -> None:
    table = (
        "Repository | Name | Category | Severity | Interactive | Status | Summary\n"
        "-----------+------+----------+----------+-------------+--------+--------\n"
        "SLES15-Updates | SUSE-SLE-2025-123 | security | important | --- | needed | Security update for openssl\n"
        "SLES15-Updates | SUSE-SLE-2025-200 | recommended | moderate | --- | needed | Recommended update for vim"
    )
    _, _, s = engine.analyze(
        json.dumps(
            linux(patch_state="manager zypper\nauto_updates 0", patches=table, updates_count="2")
        ),
        "sles",
        "15",
        NOW,
    )
    assert (s["patches"]["checked"], s["patches"]["security"]) == (True, 1)
    assert s["patches"]["items"][0]["id"] == "SUSE-SLE-2025-123"
    broken = linux(
        patch_state="manager zypper",
        patches="  Repository: Non-OSS\n  Key Expires: (EXPIRED)",
        updates_count="0",
    )
    _, _, unknown = engine.analyze(json.dumps(broken), "sles", "15", NOW)
    assert unknown["patches"]["checked"] is False  # unknown, never "nothing to patch"
    assert ids(unknown)["patch-unknown"] == "low"


def test_catalog_and_scripts() -> None:
    assert {f["key"] for f in catalog.options("linux")} >= {
        "ubuntu",
        "debian",
        "rhel",
        "centos",
        "rocky",
        "almalinux",
        "oracle",
        "amazon",
        "sles",
        "opensuse",
    }
    assert [f["key"] for f in catalog.options("windows")] == ["windows-server"]
    assert catalog.version("ubuntu", "24.04") is not None
    assert catalog.version("ubuntu", "nope") is None
    ps, sh = (SCRIPTS / "windows.ps1").read_text("utf-8"), (SCRIPTS / "linux.sh").read_text("utf-8")
    for text in (ps, sh):
        assert "READ-ONLY" in text
        assert "__SERVER_NAME__" in text
    # the scripts only read: no command that installs, removes, restarts or writes configuration
    for word in (
        "Remove-Item",
        "Set-ItemProperty",
        "Restart-",
        "Stop-Service",
        "Install-",
        "Invoke-WebRequest",
    ):
        assert word not in ps, word
    for word in (
        " rm ",
        "apt-get install",
        "apt-get upgrade",
        "dnf install",
        "systemctl restart",
        "curl ",
        "wget ",
    ):
        assert word not in sh, word


# ---- API -----------------------------------------------------------------------------------------
@pytest.fixture
def admin(client: TestClient, seeded: None) -> dict[str, str]:
    return auth_headers(client, "admin@acme.test")


def add(
    client: TestClient,
    headers: dict[str, str],
    name: str,
    platform: str = "linux",
    family: str = "ubuntu",
    version: str = "20.04",
) -> dict[str, Any]:
    r = client.post(
        "/api/v1/servers",
        json={"name": name, "platform": platform, "os_family": family, "os_version": version},
        headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()  # type: ignore[no-any-return]


def test_register_servers_get_the_script_and_save_a_check(
    client: TestClient, admin: dict[str, str]
) -> None:
    opts = client.get("/api/v1/servers/options", headers=admin).json()
    assert opts["linux"][0]["versions"][0]["label"]
    web = add(client, admin, "web01")
    win = add(client, admin, "srv-app01", "windows", "windows-server", "2016")
    assert web["os_label"] == "Ubuntu 20.04 LTS"
    assert web["latest"] is None
    # names are plain and unique per platform; the operating system must come from the list
    bad = {"platform": "linux", "os_family": "ubuntu", "os_version": "20.04"}
    assert (
        client.post(
            "/api/v1/servers", json={**bad, "name": 'x"; rm -rf /'}, headers=admin
        ).status_code
        == 422
    )
    assert (
        client.post("/api/v1/servers", json={**bad, "name": "WEB01"}, headers=admin).status_code
        == 409
    )
    assert (
        client.post(
            "/api/v1/servers",
            json={**bad, "name": "web02", "os_family": "windows-server"},
            headers=admin,
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/v1/servers", json={**bad, "name": "web02", "os_version": "1.0"}, headers=admin
        ).status_code
        == 422
    )
    # the script is tailored to the server
    sh = client.get(f"/api/v1/servers/{web['id']}/script", headers=admin).json()
    assert (sh["language"], sh["filename"]) == ("bash", "kloudrecon-health.sh")
    assert 'PKG="apt"' in sh["script"]
    assert "Server: web01" in sh["script"]
    assert "__" not in sh["script"].split("\n", 12)[2]
    ps = client.get(f"/api/v1/servers/{win['id']}/script", headers=admin).json()
    assert ps["language"] == "powershell"
    assert "server       = 'srv-app01'" in ps["script"]
    assert len(ps["steps"]) == 4
    # paste the output
    r = client.post(
        f"/api/v1/servers/{web['id']}/checks", json={"output": json.dumps(linux())}, headers=admin
    )
    assert r.status_code == 201, r.text
    detail = r.json()
    assert detail["latest"]["status"] == "critical"
    assert detail["check"]["summary"]["os"]["name"] == "Ubuntu 20.04.6 LTS"
    assert detail["check"]["summary"]["findings"]
    assert "raw" not in detail["check"]
    assert len(detail["history"]) == 1
    # rubbish and the wrong platform are refused with a clear message
    r = client.post(
        f"/api/v1/servers/{web['id']}/checks", json={"output": "not the output"}, headers=admin
    )
    assert r.status_code == 422
    assert "no JSON was found" in r.json()["error"]["message"]
    r = client.post(
        f"/api/v1/servers/{web['id']}/checks", json={"output": json.dumps(windows())}, headers=admin
    )
    assert r.status_code == 422
    assert "registered as Linux" in r.json()["error"]["message"]
    # list, fleet summary, history
    client.post(
        f"/api/v1/servers/{web['id']}/checks",
        json={"output": json.dumps(linux(oom_kills="0"))},
        headers=admin,
    )
    rows = client.get("/api/v1/servers", headers=admin, params={"platform": "linux"}).json()
    assert [(x["name"], x["checks"]) for x in rows] == [("web01", 2)]
    fleet = client.get("/api/v1/servers/summary", headers=admin).json()
    assert (fleet["servers"], fleet["checked"], fleet["by_status"]) == (2, 1, {"critical": 1})
    old = client.get(f"/api/v1/servers/{web['id']}", headers=admin).json()["history"][-1]["id"]
    assert (
        client.get(f"/api/v1/servers/{web['id']}", headers=admin, params={"check_id": old}).json()[
            "check"
        ]["id"]
        == old
    )
    assert (
        client.delete(f"/api/v1/servers/{web['id']}/checks/{old}", headers=admin).status_code == 200
    )
    # rename and delete
    r = client.patch(
        f"/api/v1/servers/{web['id']}",
        json={"name": "web01-prod", "os_version": "22.04"},
        headers=admin,
    )
    assert (r.json()["name"], r.json()["os_label"]) == ("web01-prod", "Ubuntu 22.04 LTS")
    assert client.delete(f"/api/v1/servers/{win['id']}", headers=admin).status_code == 200
    assert len(client.get("/api/v1/servers", headers=admin).json()) == 1


def test_servers_are_tenant_isolated_and_viewers_cannot_write(
    client: TestClient, admin: dict[str, str]
) -> None:
    web = add(client, admin, "web01")
    client.post(
        f"/api/v1/servers/{web['id']}/checks", json={"output": json.dumps(linux())}, headers=admin
    )
    other = register(client, "jane@contoso.test")
    owner = auth_headers(client, "owner@local.test")
    org = next(
        x
        for x in client.get("/api/v1/owner/registrations", headers=owner).json()
        if x["org"]["name"] == "Contoso Test"
    )
    client.post(
        f"/api/v1/owner/registrations/{org['org']['id']}/approve",
        json={"plan_code": "pro", "trial_days": 0},
        headers=owner,
    )
    other = auth_headers(client, "jane@contoso.test")
    assert client.get("/api/v1/servers", headers=other).json() == []
    assert client.get(f"/api/v1/servers/{web['id']}", headers=other).status_code == 404
    assert client.get(f"/api/v1/servers/{web['id']}/script", headers=other).status_code == 404
    assert (
        client.post(
            f"/api/v1/servers/{web['id']}/checks",
            json={"output": json.dumps(linux())},
            headers=other,
        ).status_code
        == 404
    )
    assert client.delete(f"/api/v1/servers/{web['id']}", headers=other).status_code == 404
    assert client.get("/api/v1/servers").status_code == 401


def test_server_health_report_is_a_word_document(client: TestClient, admin: dict[str, str]) -> None:
    from docx import Document

    # no checked server yet: nothing to report on
    r = client.post("/api/v1/reports", json={"type": "server_health"}, headers=admin)
    assert r.status_code == 422
    assert "No server has a health check yet" in r.json()["error"]["message"]
    web = add(client, admin, "web01")
    win = add(client, admin, "srv-app01", "windows", "windows-server", "2016")
    add(client, admin, "never-checked")
    client.post(
        f"/api/v1/servers/{web['id']}/checks", json={"output": json.dumps(linux())}, headers=admin
    )
    client.post(
        f"/api/v1/servers/{win['id']}/checks", json={"output": json.dumps(windows())}, headers=admin
    )
    r = client.post(
        "/api/v1/reports",
        json={"type": "server_health", "approved_by": "Dana Reyes"},
        headers=admin,
    )
    assert r.status_code == 202, r.text
    assert (r.json()["format"], r.json()["title"], r.json()["scan_id"]) == (
        "docx",
        "Server Health Report",
        None,
    )
    assert run_jobs("report.render") == 1
    report = client.get(f"/api/v1/reports/{r.json()['id']}", headers=admin).json()
    assert report["status"] == "ready", report["error"]
    link = client.post(f"/api/v1/reports/{report['id']}/link", headers=admin).json()
    got = client.get(link["url"].split("8000", 1)[-1])
    assert got.headers["content-disposition"].endswith('.docx"')
    d = Document(io.BytesIO(got.content))
    h1 = [p.text.split("  ", 1)[1] for p in d.paragraphs if p.style.name == "Heading 1"]
    assert h1 == [
        "Executive summary",
        "Scope and method",
        "Findings across all servers",
        "Server: srv-app01",
        "Server: web01",
        "Recommendations",
        "Appendix: glossary",
    ]
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts.extend(cell.text for row in t.rows for cell in row.cells)
    text = " | ".join(parts)
    for expected in (
        "Document control",
        "Dana Reyes",
        "Servers checked",
        "SMBv1 is enabled",
        "CVE-2026-1111",
        "Disk / is 97% full",
        "KloudRecon never connected to the servers",
        "Security configuration",
        "Listening ports",
    ):
        assert expected in text, expected
    assert "never-checked" not in text  # only servers with a check are reported
