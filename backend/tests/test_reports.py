"""PDF reports on the Contoso Demo: all 7 types render, branding / white-label / watermark,
15-minute signed links, token-protected print + download, schedules with email delivery."""

from __future__ import annotations

import base64
import time
import uuid
from collections.abc import Iterator
from datetime import UTC, datetime
from typing import Any

import psycopg
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from kloudrecon.core import signing
from kloudrecon.core.db import session_scope
from kloudrecon.core.rls import enable_system, set_tenant
from kloudrecon.domain.report_models import REPORT_TYPES, WORD_TYPES
from kloudrecon.reports import charts
from kloudrecon.reports.document import Branding, ReportDoc
from tests.conftest import TABLES, _admin_dsn, auth_headers, run_jobs

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


@pytest.fixture(autouse=True)
def _clean_tables() -> None:  # data is shared across this module (see `demo`)
    return None


@pytest.fixture(scope="module")
def demo(_database: None) -> Iterator[dict[str, Any]]:
    from kloudrecon.services.demo import seed_demo_data
    from kloudrecon.services.seed import seed_dev

    with psycopg.connect(_admin_dsn(), autocommit=True) as conn:
        conn.execute(f"TRUNCATE {TABLES} RESTART IDENTITY CASCADE")
    with session_scope() as s:
        seed_dev(s)
    result = seed_demo_data(history_days=(2.0, 1.0))
    yield {"org_id": result.org_id}
    with psycopg.connect(_admin_dsn(), autocommit=True) as conn:
        conn.execute(f"TRUNCATE {TABLES} RESTART IDENTITY CASCADE")


@pytest.fixture(scope="module")
def api(demo: dict[str, Any]) -> Iterator[TestClient]:
    from kloudrecon.api.main import create_app

    with TestClient(create_app(start_listener=False)) as c:
        yield c


def _admin(api: TestClient) -> dict[str, str]:
    return auth_headers(api, "demo@contoso.test")


def _run(job_type: str) -> int:
    return run_jobs(job_type)


# ---- building blocks -----------------------------------------------------------------------------
def test_charts_are_escaped_svg() -> None:
    svg = charts.bar_chart(["<b>jan</b>", "feb"], [10.0, 20.0], title="t<script>")
    assert svg.startswith("<svg")
    assert "<script>" not in svg
    assert "&lt;b&gt;jan" in svg
    assert charts.donut([("a", 1, None), ("b", 3, None)], title="d").count("<circle") == 3
    assert "—" in charts.score_ring(None, label="Score")
    assert charts.bar_chart([], [], title="x") == ""


def test_document_escapes_customer_text() -> None:
    doc = ReportDoc(
        title="Report",
        subtitle="s",
        branding=Branding(company="<img src=x onerror=alert(1)>"),
        meta=[],
    )
    doc.section("One").table(["Name"], [["</td><script>alert(1)</script>"]])
    html = doc.render()
    assert "<script>alert(1)" not in html
    assert "&lt;img src=x" in html
    assert "window.__reportReady = true" in html


# ---- the audit report is a Word document (decision 0053) ---------------------------------------
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
AUDIT_SECTIONS = [
    "Executive summary",
    "General information",
    "Environment overview",
    "Security findings",
    "Compliance",
    "Reliability and governance",
    "Cost analysis",
    "Recommendations",
    "Financial summary",
    "Appendix A: findings register",
    "Appendix B: resource inventory",
    "Appendix C: scope and method",
    "Appendix D: glossary",
]


def test_audit_report_is_a_word_document_in_the_agreed_order(api: TestClient) -> None:
    import io

    from docx import Document

    body = {"type": "audit", "prepared_by": "Alex Admin", "approved_by": "Dana Reyes"}
    resp = api.post("/api/v1/reports", json=body, headers=_admin(api))
    assert resp.status_code == 202, resp.text
    rid = resp.json()["id"]
    assert resp.json()["format"] == "docx"
    assert _run("report.render") == 1
    report = api.get(f"/api/v1/reports/{rid}", headers=_admin(api)).json()
    assert report["status"] == "ready", report["error"]
    assert report["format"] == "docx"
    assert report["title"] == "Azure Audit Report"
    assert report["pages"] is None  # Word lays the pages out
    assert report["params"] == {"prepared_by": "Alex Admin", "approved_by": "Dana Reyes"}
    link = api.post(f"/api/v1/reports/{rid}/link", headers=_admin(api)).json()
    got = api.get(link["url"].split("8000", 1)[-1])
    assert got.status_code == 200
    assert got.headers["content-type"] == DOCX_MIME
    assert got.headers["content-disposition"].endswith('.docx"')
    d = Document(io.BytesIO(got.content))
    h1 = [p.text.split("  ", 1)[1] for p in d.paragraphs if p.style.name == "Heading 1"]
    assert h1 == AUDIT_SECTIONS
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts.extend(cell.text for row in t.rows for cell in row.cells)
    text = " | ".join(parts)
    # page 2: approval block with the names given, a blank line where none was given
    for label in ("Document control", "Prepared by", "Reviewed by", "Report approved by"):
        assert label in text
    assert "Alex Admin" in text
    assert "Dana Reyes" in text
    # general information, findings, recommendations with an owner column
    assert "Directory (tenant) ID" in text
    assert "Limitations" in text
    assert "vm-prod-jumpbox01" in text
    assert "Patching and software lifecycle" in text
    assert "Act now" in text
    assert "Owner" in text
    # the financial summary comes after the recommendations
    assert "Spend after savings" in text
    assert "Savings by category" in text
    assert "Assumptions" in text
    assert h1.index("Financial summary") == h1.index("Recommendations") + 1
    # compliance wording
    assert "not a certification" in text
    # the print view still works, and deleting removes the Word file too
    assert api.get(link["print_url"].split("8000", 1)[-1]).status_code == 200
    assert api.delete(f"/api/v1/reports/{rid}", headers=_admin(api)).status_code == 200


# ---- rendering -----------------------------------------------------------------------------------
@pytest.mark.parametrize("report_type", [t for t in REPORT_TYPES if t not in WORD_TYPES])
def test_every_report_type_renders(api: TestClient, report_type: str) -> None:
    body: dict[str, Any] = {"type": report_type}
    if report_type == "compliance":
        body["framework"] = "cis"
    if report_type == "changes":
        body["days"] = 30
    resp = api.post("/api/v1/reports", json=body, headers=_admin(api))
    assert resp.status_code == 202, resp.text
    rid = resp.json()["id"]
    assert resp.json()["status"] == "queued"
    assert _run("report.render") == 1
    report = api.get(f"/api/v1/reports/{rid}", headers=_admin(api)).json()
    assert report["status"] == "ready", report["error"]
    assert report["pages"] >= 1
    assert report["size_bytes"] > 1000
    link = api.post(f"/api/v1/reports/{rid}/link", headers=_admin(api)).json()
    pdf = api.get(link["url"].split("8000", 1)[-1])
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert "attachment" in pdf.headers["content-disposition"]
    html = api.get(link["print_url"].split("8000", 1)[-1])
    assert html.status_code == 200
    assert "<!doctype html>" in html.text
    assert "Contoso" in html.text
    assert "default-src 'none'" in html.headers["content-security-policy"]


def test_full_report_has_every_chapter(api: TestClient) -> None:
    rid = api.post("/api/v1/reports", json={"type": "full"}, headers=_admin(api)).json()["id"]
    _run("report.render")
    link = api.post(f"/api/v1/reports/{rid}/link", headers=_admin(api)).json()
    html = api.get(link["print_url"].split("8000", 1)[-1]).text
    for chapter in (
        "Executive summary",
        "Inventory",
        "Security posture",
        "Compliance",
        "Reliability",
        "Governance",
        "Cost &amp; FinOps",
        "Architecture",
        "Recommendations backlog",
        "Appendix",
        "Contents",
    ):
        assert chapter in html, chapter
    assert "certif" in html  # "not a certification" disclaimer
    assert "<svg" in html


def test_links_are_signed_and_expire(api: TestClient) -> None:
    rid = api.post("/api/v1/reports", json={"type": "executive"}, headers=_admin(api)).json()["id"]
    _run("report.render")
    assert api.get(f"/api/v1/public/reports/{rid}/download?token={'x' * 40}").status_code == 403
    other = signing.sign("report", str(uuid.uuid4()), ttl_seconds=60)
    assert api.get(f"/api/v1/public/reports/{rid}/download?token={other}").status_code == 403
    expired = signing.sign("report", rid, ttl_seconds=-1)
    resp = api.get(f"/api/v1/public/reports/{rid}/download?token={expired}")
    assert resp.status_code == 403
    assert resp.json()["error"]["code"] == "link_expired"
    good = signing.sign("report", rid, ttl_seconds=60)
    assert api.get(f"/api/v1/public/reports/{rid}/download?token={good}").status_code == 200
    body = signing.verify(good, "report", rid)
    assert body["e"] <= int(time.time()) + 60


def test_reports_are_tenant_isolated(api: TestClient) -> None:
    rid = api.get("/api/v1/reports", headers=_admin(api)).json()[0]["id"]
    acme = auth_headers(api, "admin@acme.test")
    assert api.get(f"/api/v1/reports/{rid}", headers=acme).status_code == 404
    assert api.post(f"/api/v1/reports/{rid}/link", headers=acme).status_code == 404
    assert api.get("/api/v1/reports", headers=acme).json() == []


def test_compliance_pack_requires_framework(api: TestClient) -> None:
    resp = api.post("/api/v1/reports", json={"type": "compliance"}, headers=_admin(api))
    assert resp.status_code == 422


# ---- branding ------------------------------------------------------------------------------------
def test_branding_white_label_and_logo(api: TestClient) -> None:
    headers = _admin(api)
    logo = "data:image/png;base64," + base64.b64encode(PNG_1PX).decode()
    resp = api.put(
        "/api/v1/settings/branding",
        json={
            "primary_color": "#7C3AED",
            "prepared_by": "Fabrikam Consulting",
            "logo_data_url": logo,
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["white_label"] is True
    assert out["has_logo"] is True
    assert out["primary_color"] == "#7C3AED"
    rid = api.post("/api/v1/reports", json={"type": "executive"}, headers=headers).json()["id"]
    _run("report.render")
    link = api.post(f"/api/v1/reports/{rid}/link", headers=headers).json()
    html = api.get(link["print_url"].split("8000", 1)[-1]).text
    assert "#7C3AED" in html
    assert "Fabrikam Consulting" in html
    assert "data:image/png;base64," in html
    assert 'class="mark"' not in html  # KloudRecon mark removed (white-label)
    bad = api.put(
        "/api/v1/settings/branding",
        json={
            "logo_data_url": "data:image/png;base64," + base64.b64encode(b"<svg onload=x>").decode()
        },
        headers=headers,
    )
    assert bad.status_code == 422
    assert (
        api.put(
            "/api/v1/settings/branding", json={"primary_color": "red;}"}, headers=headers
        ).status_code
        == 422
    )
    api.put("/api/v1/settings/branding", json={"remove_logo": True}, headers=headers)


def test_trial_reports_have_a_watermark_and_no_white_label(
    api: TestClient, demo: dict[str, Any]
) -> None:
    with psycopg.connect(_admin_dsn(), autocommit=True) as conn:
        pro = conn.execute(
            "SELECT plan_id FROM organizations WHERE id = %s", (demo["org_id"],)
        ).fetchone()
        conn.execute(
            "UPDATE organizations SET plan_id = (SELECT id FROM plans WHERE code='trial') "
            "WHERE id = %s",
            (demo["org_id"],),
        )
    try:
        headers = _admin(api)
        resp = api.put(
            "/api/v1/settings/branding", json={"primary_color": "#000000"}, headers=headers
        )
        assert resp.status_code == 403
        assert resp.json()["error"]["code"] == "plan_feature"
        rid = api.post("/api/v1/reports", json={"type": "executive"}, headers=headers).json()["id"]
        _run("report.render")
        link = api.post(f"/api/v1/reports/{rid}/link", headers=headers).json()
        html = api.get(link["print_url"].split("8000", 1)[-1]).text
        assert 'class="watermark">TRIAL' in html
        sched = api.post(
            "/api/v1/reports/schedules",
            json={"type": "executive", "frequency": "weekly", "recipients": ["demo@contoso.test"]},
            headers=headers,
        )
        assert sched.status_code == 403
    finally:
        with psycopg.connect(_admin_dsn(), autocommit=True) as conn:
            conn.execute(
                "UPDATE organizations SET plan_id = %s WHERE id = %s",
                (pro[0] if pro else None, demo["org_id"]),
            )


# ---- schedules -----------------------------------------------------------------------------------
def test_next_run() -> None:
    from kloudrecon.domain.report_schemas import ScheduleIn
    from kloudrecon.services.reports import next_run

    weekly = ScheduleIn(
        type="executive", frequency="weekly", weekday=0, hour_utc=6, recipients=["a@b.test"]
    )
    wed = datetime(2026, 9, 30, 10, 0, tzinfo=UTC)
    assert next_run(weekly, wed) == datetime(2026, 10, 5, 6, 0, tzinfo=UTC)
    monthly = ScheduleIn(
        type="finops", frequency="monthly", day_of_month=1, hour_utc=7, recipients=["a@b.test"]
    )
    assert next_run(monthly, wed) == datetime(2026, 10, 1, 7, 0, tzinfo=UTC)
    dec = datetime(2026, 12, 15, tzinfo=UTC)
    assert next_run(monthly, dec) == datetime(2027, 1, 1, 7, 0, tzinfo=UTC)


def test_schedule_delivers_pdf_by_email(api: TestClient, demo: dict[str, Any]) -> None:
    from kloudrecon.services import reports
    from kloudrecon.services.infra import get_email_sender

    headers = _admin(api)
    outside = api.post(
        "/api/v1/reports/schedules",
        json={"type": "finops", "frequency": "monthly", "recipients": ["stranger@evil.test"]},
        headers=headers,
    )
    assert outside.status_code == 422
    resp = api.post(
        "/api/v1/reports/schedules",
        json={
            "type": "finops",
            "frequency": "monthly",
            "day_of_month": 1,
            "recipients": ["demo@contoso.test"],
        },
        headers=headers,
    )
    assert resp.status_code == 201, resp.text
    sid = resp.json()["id"]
    viewer = auth_headers(api, "viewer@contoso.test")
    assert api.get("/api/v1/reports/schedules", headers=viewer).status_code == 200
    assert api.delete(f"/api/v1/reports/schedules/{sid}", headers=viewer).status_code == 403
    with session_scope() as s:
        enable_system(s)
        s.execute(text("UPDATE report_schedules SET next_run_at = now() - interval '1 minute'"))
        assert reports.run_due_schedules(s) == 1
        assert reports.run_due_schedules(s) == 0  # moved to the next month
    sender = get_email_sender()
    before = len(sender.sent)  # type: ignore[attr-defined]
    assert _run("report.render") == 1
    mails = sender.sent[before:]  # type: ignore[attr-defined]
    delivered = [m for m in mails if m.tags.get("template") == "report_ready"]
    assert delivered
    assert delivered[0].to == ["demo@contoso.test"]
    name, ctype, data = delivered[0].attachments[0]
    assert name.endswith(".pdf")
    assert ctype == "application/pdf"
    assert data.startswith(b"%PDF")
    with session_scope() as s:
        set_tenant(s, demo["org_id"])
        assert (
            s.execute(text("SELECT count(*) FROM reports WHERE schedule_id IS NOT NULL")).scalar()
            == 1
        )
